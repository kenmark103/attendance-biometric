"""Synthetic Zoho leave-report generator for tests.

Writes workbooks in the exact HR layout with FAKE IDs/names only — never
derived from the real report. Outputs go to a temp dir (never committed).

Layouts: genuine .xls (xlwt/OLE2), .xlsx (openpyxl/ZIP), HTML-disguised
.xls (stdlib). Edge-case switches mirror the ingestion spec test plan.

Usage:
    python scripts/make_leave_fixture.py --out <dir> [--months 2]
        [--employees 6] [--format xlsx|xls|html] [--edge all|...]
"""
from __future__ import annotations

import argparse
import calendar as calmod
from datetime import date

DEPT = "GBS- ENGINEERING CENTRE,Engineering Centre"

WEEKDAY_BASE = "P"


def _month_days(year: int, month: int) -> list[date]:
    ndays = calmod.monthrange(year, month)[1]
    return [date(year, month, d) for d in range(1, ndays + 1)]


def _base_code(day: date) -> str:
    return "W" if day.weekday() >= 5 else WEEKDAY_BASE


def build_grid(year: int, month: int, n_emp: int, edges: set[str],
               sheet_index: int, missing_a1: bool = False) -> list[list[object]]:
    """Build the raw cell grid for one month sheet."""
    days = _month_days(year, month)
    mon = days[0].strftime("%b")
    rows: list[list[object]] = []
    rows.append(["Department", DEPT])
    rows.append([])
    rows.append(["Start Date", days[0].strftime("%d-%b-%Y")])
    rows.append(["End Date", days[-1].strftime("%d-%b-%Y")])
    rows.append(["", "", "P", " - Present", "A", " - Absent"])
    rows.append(["", "", "EL", " - Annual Leave", "SL", " - Sick Leave"])
    rows.append(["Employee Id", "Employee Name"] + [f"{d.day} - {mon}" for d in days])
    if missing_a1:
        rows[0][0] = ""
    for n in range(1, n_emp + 1):
        emp_id = f"{90000 + n}"
        # Leaver switch: last employee vanishes from later sheets.
        if "leaver" in edges and sheet_index > 0 and n == n_emp:
            continue
        codes = [_base_code(d) for d in days]
        if "leave_run" in edges and n == 1:
            # 3-day EL run on first three weekdays.
            wed = [i for i, d in enumerate(days) if d.weekday() < 5][:3]
            for i in wed:
                codes[i] = "EL"
        if "compound" in edges and n == 2:
            wd = next(i for i, d in enumerate(days) if d.weekday() < 5)
            codes[wd] = "EL/P"
        if "half" in edges and n == 2:
            wd = next(i for i, d in enumerate(days) if d.weekday() < 5 and codes[i] == "P")
            codes[wd] = "0.5EL/0.5P"
        if "interior_dash" in edges and n == 3:
            mid = len(days) // 2
            codes[mid] = "-"
        if "unknown" in edges and n == 3:
            wd = next(i for i, d in enumerate(days) if d.weekday() < 5 and codes[i] == "P")
            codes[wd] = "CL"
        if "provisional" in edges and sheet_index == 1:
            # Final column mostly dashes (unprocessed day).
            codes[-1] = "-"
        rows.append([emp_id, f"Test User {n}"] + codes)
    if "duplicate_identical" in edges:
        rows.append(list(rows[7]))
    if "duplicate_conflict" in edges:
        dup = list(rows[7])
        last = len(dup) - 1
        dup[last] = "EL" if dup[last] != "EL" else "SL"
        rows.append(dup)
    return rows


def write_xlsx(path: str, sheets: list[tuple[str, list[list[object]]]]) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, grid in sheets:
        ws = wb.create_sheet(title=name)
        for row in grid:
            ws.append(row)
    wb.save(path)


def write_xls(path: str, sheets: list[tuple[str, list[list[object]]]]) -> None:
    import xlwt

    wb = xlwt.Workbook()
    for name, grid in sheets:
        ws = wb.add_sheet(name[:31])
        for r, row in enumerate(grid):
            for c, val in enumerate(row):
                ws.write(r, c, "" if val is None else val)
    wb.save(path)


def write_html_xls(path: str, sheets: list[tuple[str, list[list[object]]]]) -> None:
    def esc(v: object) -> str:
        s = "" if v is None else str(v)
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    parts = ["<html><body>"]
    for _name, grid in sheets:
        parts.append("<table>")
        for row in grid:
            parts.append("<tr>" + "".join(f"<td>{esc(v)}</td>" for v in row) + "</tr>")
        parts.append("</table>")
    parts.append("</body></html>")
    with open(path, "w", encoding="utf-8") as f:
        f.write("".join(parts))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--months", type=int, default=2)
    ap.add_argument("--employees", type=int, default=6)
    ap.add_argument("--format", default="xlsx", choices=["xlsx", "xls", "html"])
    ap.add_argument("--edge", action="append", default=[],
                    help="repeatable; use 'all' for every edge case")
    args = ap.parse_args()

    edges = set(args.edge)
    if "all" in edges:
        edges = {"leave_run", "compound", "half", "interior_dash", "unknown",
                 "duplicate_identical", "leaver", "provisional", "missing_a1"}
    year, start_month = 2026, 1
    sheets = []
    for i in range(args.months):
        m = start_month + i
        grid = build_grid(year, m, args.employees, edges, i,
                          missing_a1=("missing_a1" in edges and i == 0))
        sheets.append((f"M{i + 1}", grid))
    import os

    os.makedirs(args.out, exist_ok=True)
    ext = {"xlsx": "xlsx", "xls": "xls", "html": "xls"}[args.format]
    path = os.path.join(args.out, f"leave_fixture.{ext}")
    if args.format == "xlsx":
        write_xlsx(path, sheets)
    elif args.format == "xls":
        write_xls(path, sheets)
    else:
        write_html_xls(path, sheets)
    print(f"wrote {path}")


if __name__ == "__main__":
    main()
